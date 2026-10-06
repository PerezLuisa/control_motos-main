"""
Lectura del archivo de carga diario «REFRENCIAS MOTOS AKT.xlsx» (también acepta CSV).

Estructura (siempre la misma):
    Artículo | Marca | Descripción | COD INT
    7700149546881 | AKT | Moto AK125TTR CBS G/N/V 26 PT | AK125TTR EIII

Una fila = una moto. Si un artículo se repite, son varias unidades de esa referencia.
Los nombres de columna se reconocen sin importar mayúsculas, tildes ni espacios.
El código de la carga, la fecha, el destino, la placa y el conductor se escriben en el formulario.
"""
import csv
import io
import unicodedata

NOMBRE_ARCHIVO = "REFRENCIAS MOTOS AKT.xlsx"
COLUMNAS = {  # campo -> (nombre en el archivo, sinónimos aceptados)
    "articulo": ("Artículo", ["articulo", "codigo_articulo", "cod_articulo", "codigo", "ean", "sku"]),
    "marca": ("Marca", ["marca"]),
    "descripcion": ("Descripción", ["descripcion", "descripcion_articulo"]),
    "referencia": ("COD INT", ["cod_int", "codigo_interno", "cod_interno"]),
}


class ErrorImportacion(Exception):
    pass


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = texto.strip().lower()
    for c in " .-/º°#":
        texto = texto.replace(c, "_")
    while "__" in texto:
        texto = texto.replace("__", "_")
    return texto.strip("_")


def _mapear_encabezado(encabezado):
    """{campo: índice_de_columna} para los encabezados reconocidos."""
    normalizados = [_normalizar(h) for h in encabezado]
    mapa = {}
    for campo, (_, alias) in COLUMNAS.items():
        for i, h in enumerate(normalizados):
            if h in alias:
                mapa[campo] = i
                break
    return mapa


def _filas_csv(contenido):
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    return list(csv.reader(io.StringIO(texto), dialecto))


def _filas_xlsx(contenido):
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    filas = [list(f) for f in libro.worksheets[0].iter_rows(values_only=True)]
    libro.close()
    return filas


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # códigos que Excel guarda como número
    return str(valor).strip()


def leer_archivo(nombre_archivo, contenido):
    """Devuelve la lista de motos: [{articulo, marca, descripcion, referencia}]."""
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        filas = _filas_xlsx(contenido)
    elif nombre.endswith((".csv", ".txt")):
        filas = _filas_csv(contenido)
    else:
        raise ErrorImportacion(f"Formato no soportado. Suba el archivo «{NOMBRE_ARCHIVO}».")

    # El encabezado es la primera fila (de las 10 primeras) con la columna Artículo
    for idx, fila in enumerate(filas[:10]):
        mapa = _mapear_encabezado(fila)
        if "articulo" in mapa:
            break
    else:
        raise ErrorImportacion(
            f"El archivo no tiene la estructura de «{NOMBRE_ARCHIVO}»: falta la columna Artículo. "
            "Columnas esperadas: Artículo, Marca, Descripción, COD INT."
        )
    faltan = [nombre for campo, (nombre, _) in COLUMNAS.items() if campo not in mapa]
    if faltan:
        raise ErrorImportacion(f"El archivo no tiene la estructura de «{NOMBRE_ARCHIVO}»: faltan las columnas {', '.join(faltan)}.")

    motos, errores = [], []
    for n, fila in enumerate(filas[idx + 1:], start=idx + 2):
        valor = lambda campo: _texto(fila[mapa[campo]]) if mapa[campo] < len(fila) else ""  # noqa: E731
        if not any(_texto(c) for c in fila):
            continue  # fila vacía
        articulo = valor("articulo").replace(" ", "")
        if not articulo:
            errores.append(f"Fila {n}: no tiene código de artículo.")
            continue
        motos.append({
            "articulo": articulo,
            "marca": valor("marca") or None,
            "descripcion": valor("descripcion") or None,
            "referencia": valor("referencia") or None,
        })

    if errores:
        raise ErrorImportacion("El archivo tiene errores:\n" + "\n".join(errores[:20]))
    if not motos:
        raise ErrorImportacion("El archivo no tiene motos.")
    return motos
